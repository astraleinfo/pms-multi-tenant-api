"""Onboarding a clinic from a spreadsheet: its staff, and its patients.

The console downloads a sample sheet, the clinic fills it in, the console
uploads it back. Three rules keep that safe:

*   **A preview first.** Every upload is checked row by row and reported back
    before anything is written, with the reason each bad row is bad.
*   **All or nothing.** An import commits only when every row is valid. Half a
    staff list is worse than none: the people left out do not know they were.
*   **The same rules as the forms.** A patient row goes through the clinic's
    own effective schema — its custom fields and which of them are required —
    and a staff row through the same checks as "add a person". A spreadsheet
    is not a side door around validation.

Staff are created without a password and emailed an invitation, exactly as
if an admin had added them one by one. Nobody's password travels in a sheet.

`.xlsx` is what the sample sheets are; `.csv` is accepted too, because it is
what most other systems export.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from email_validator import EmailNotValidError, validate_email
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import ValidationFailed
from app.core.permissions import Role
from app.models.clinical import Patient
from app.models.enums import FieldEntity, FieldType, Sex
from app.models.identity import User
from app.models.tenancy import FieldDefinition, Tenant
from app.schemas.clinical import PatientIn
from app.services import tenants as tenant_service

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = {"users": 300, "patients": 3000}

ROLE_LABELS = {
    Role.TENANT_ADMIN.value: "Practice admin",
    Role.DOCTOR.value: "Clinician",
    Role.RECEPTIONIST.value: "Receptionist",
    Role.PHARMACIST.value: "Pharmacist",
    Role.LAB_TECH.value: "Lab technician",
}


#: How the choices are named back to a person: the words first, the code the
#: sample sheet's dropdown uses second, so either can be matched to the sheet.
_ROLE_CHOICES = ", ".join(f"{label} ({code})" for code, label in ROLE_LABELS.items())


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    required: bool = False
    help: str = ""
    example: Any = ""
    #: a fixed list the sheet offers as a dropdown
    options: tuple[str, ...] = ()
    #: for a custom field, the definition it is validated against
    definition: FieldDefinition | None = None


USER_COLUMNS: tuple[Column, ...] = (
    Column("name", "Full name", True, "As it should appear on screens and letters.", "Dr Meera Rao"),
    Column("email", "Email", True, "Their own address — the invitation goes here. Unique in the clinic.",
           "meera.rao@example.com"),
    Column("role", "Role", True,
           "One of: " + ", ".join(f"{label} ({code})" for code, label in ROLE_LABELS.items())
           + ". Either the name or the code in brackets works.",
           "doctor", tuple(ROLE_LABELS)),
    Column("speciality", "Speciality", False, "For clinicians: shown next to their name.", "Endodontics"),
)

_PATIENT_CORE: tuple[Column, ...] = (
    Column("full_name", "Full name", True, "", "Anita Sharma"),
    Column("dob", "Date of birth", True, "YYYY-MM-DD, or a date cell. DD/MM/YYYY also works.", "1986-04-12"),
    Column("sex", "Sex", False, "female, male or other. Blank means other.", "female",
           tuple(s.value for s in Sex)),
    Column("phone", "Phone", False, "With country code for SMS reminders.", "+91 98450 12345"),
    Column("email", "Email", False, "", "anita.sharma@example.com"),
    Column("address", "Address", False, "", "14 MG Road, Bengaluru"),
    Column("blood_group", "Blood group", False, "e.g. O+, AB-", "B+"),
    Column("allergies", "Allergies", False, "Blank means none recorded.", "Penicillin"),
)


def patient_columns(schema: list[FieldDefinition]) -> tuple[Column, ...]:
    """The built-in patient fields, then this clinic's own — each with the rule
    the form applies to it."""
    built_in = {c.key for c in _PATIENT_CORE}
    extra = []
    for d in sorted(schema, key=lambda d: d.order):
        if d.field_key in built_in:
            continue
        kind = FieldType(d.type)
        help_text = d.help or ""
        if kind == FieldType.DATE:
            help_text = (help_text + " YYYY-MM-DD.").strip()
        elif kind == FieldType.BOOLEAN:
            help_text = (help_text + " yes or no.").strip()
        elif kind == FieldType.MULTISELECT:
            help_text = (help_text + " Separate choices with ;").strip()
        elif kind == FieldType.NUMBER:
            help_text = (help_text + " A number.").strip()
        options = tuple(d.options or ()) if kind == FieldType.SELECT else ()
        if kind == FieldType.BOOLEAN:
            options = ("yes", "no")
        extra.append(Column(d.field_key, d.label, bool(d.required), help_text, "", options, d))
    return _PATIENT_CORE + tuple(extra)


# --- the sample sheet ----------------------------------------------------------------


def build_template(kind: str, columns: tuple[Column, ...], *, clinic: str) -> bytes:
    """An .xlsx with the columns in row 1, one worked example in row 2, a
    dropdown on every column with fixed choices, and a sheet explaining each."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Staff" if kind == "users" else "Patients"

    head_fill = PatternFill("solid", fgColor="2F6BF3")
    required_fill = PatternFill("solid", fgColor="1A4FCC")
    example_font = Font(italic=True, color="6B7890")

    for i, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=i, value=f"{column.label}{' *' if column.required else ''}")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if column.required else head_fill
        cell.alignment = Alignment(vertical="center")
        if column.help:
            cell.comment = Comment(column.help, "Astrale")
        example = sheet.cell(row=2, column=i, value=column.example)
        example.font = example_font
        width = max(len(column.label) + 4, len(str(column.example)) + 2, 14)
        sheet.column_dimensions[cell.column_letter].width = min(width, 40)

        if column.options:
            listing = '"' + ",".join(column.options) + '"'
            if len(listing) <= 255:  # Excel's own limit on an inline list
                rule = DataValidation(type="list", formula1=listing, allow_blank=not column.required)
                rule.error = f"Choose one of: {', '.join(column.options)}"
                rule.errorTitle = column.label
                sheet.add_data_validation(rule)
                rule.add(f"{cell.column_letter}2:{cell.column_letter}{MAX_ROWS[kind] + 1}")

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22

    guide = book.create_sheet("How to fill this in")
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 80
    what = "staff accounts" if kind == "users" else "patient records"
    lines = [
        (f"Import {what} into {clinic}", "", ""),
        ("", "", ""),
        ("Replace row 2 (the grey example) with your own, one per row, on the first sheet.", "", ""),
        ("Columns marked * are required. Keep the header row as it is.", "", ""),
        (f"Up to {MAX_ROWS[kind]} rows per file. Nothing is saved until every row passes.", "", ""),
    ]
    if kind == "users":
        lines.append(("Each person is emailed an invitation to choose their own password.", "", ""))
    lines += [("", "", ""), ("Column", "Required", "What goes in it")]
    lines += [(c.label, "yes" if c.required else "", c.help or "") for c in columns]
    for row in lines:
        guide.append(row)
    guide["A1"].font = Font(bold=True, size=14)
    header_row = len(lines) - len(columns)
    for cell in guide[header_row]:
        cell.font = Font(bold=True)

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# --- reading an upload ------------------------------------------------------------------


def _normalise(header: Any) -> str:
    text = re.sub(r"[*()]", "", str(header or "")).strip().lower()
    return re.sub(r"[\s_-]+", " ", text)


def read_rows(
    content: bytes, filename: str, columns: tuple[Column, ...], *, kind: str
) -> list[tuple[int, dict[str, Any]]]:
    """(sheet row number, {column key: raw value}) for every non-empty row.

    Headers are matched by label or key, in any order and any case, so a sheet
    somebody re-typed the headers of still reads. A missing required column is
    one clear error rather than the same error on every row.
    """
    if not content:
        raise ValidationFailed("That file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ValidationFailed(f"Files are limited to {MAX_FILE_BYTES // (1024 * 1024)} MB.")

    name = (filename or "").lower()
    if name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValidationFailed("That CSV is not UTF-8 text. Save it as 'CSV UTF-8'.") from exc
        grid = list(csv.reader(io.StringIO(text)))
    elif name.endswith(".xlsx"):
        grid = _xlsx_grid(content)
    else:
        raise ValidationFailed("Upload the .xlsx sample sheet, or a .csv with the same columns.")

    if not grid:
        raise ValidationFailed("That sheet has no header row.")

    lookup: dict[str, Column] = {}
    for column in columns:
        lookup[_normalise(column.label)] = column
        lookup[_normalise(column.key)] = column
    index: dict[int, Column] = {}
    for i, header in enumerate(grid[0]):
        column = lookup.get(_normalise(header))
        if column is not None and column not in index.values():
            index[i] = column

    missing = [c.label for c in columns if c.required and c not in index.values()]
    if missing:
        raise ValidationFailed(
            "The sheet is missing required columns: " + ", ".join(missing) + ".",
            details={"missing_columns": missing},
        )

    rows: list[tuple[int, dict[str, Any]]] = []
    for number, raw in enumerate(grid[1:], start=2):
        values = {col.key: (raw[i] if i < len(raw) else None) for i, col in index.items()}
        if all(_blank(v) for v in values.values()):
            continue
        rows.append((number, values))

    if not rows:
        raise ValidationFailed("The sheet has a header row but no rows under it.")
    limit = MAX_ROWS[kind]
    if len(rows) > limit:
        raise ValidationFailed(f"Up to {limit} rows per file — this one has {len(rows)}. Split it.")
    return rows


def _xlsx_grid(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable file is the same answer
        raise ValidationFailed("That file could not be read as an .xlsx workbook.") from exc
    try:
        sheet = book.worksheets[0]
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        book.close()
    while grid and all(_blank(v) for v in grid[-1]):
        grid.pop()
    return grid


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("not a date")


# --- checking rows ------------------------------------------------------------------------


@dataclass
class RowResult:
    row: int
    label: str
    errors: list[str] = field(default_factory=list)
    data: dict[str, Any] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors


async def check_users(
    session: AsyncSession, tenant: Tenant, rows: list[tuple[int, dict[str, Any]]]
) -> list[RowResult]:
    """The caller has the clinic bound."""
    existing = set(
        (
            await session.execute(
                select(func.lower(User.email)).where(User.tenant_id == tenant.id)
            )
        ).scalars()
    )
    by_label = {label.lower(): role for role, label in ROLE_LABELS.items()}
    seen: dict[str, int] = {}
    results = []
    for number, values in rows:
        name = _text(values.get("name"))
        result = RowResult(row=number, label=name or "(no name)")
        if len(name) < 2:
            result.errors.append("Full name is required.")
        elif len(name) > 160:
            result.errors.append("Full name is too long.")

        email = _text(values.get("email")).lower()
        if not email:
            result.errors.append("Email is required.")
        else:
            try:
                email = validate_email(email, check_deliverability=False).normalized.lower()
            except EmailNotValidError:
                result.errors.append(f"'{email}' is not an email address.")
            if email in existing:
                result.errors.append(f"{email} already has an account in this clinic.")
            elif email in seen:
                result.errors.append(f"{email} is also on row {seen[email]}.")
            seen.setdefault(email, number)

        raw_role = _text(values.get("role")).lower().replace(" ", "_")
        role = raw_role if raw_role in ROLE_LABELS else by_label.get(_text(values.get("role")).lower())
        if not role:
            got = _text(values.get("role"))
            result.errors.append(
                f"Role must be one of {_ROLE_CHOICES}"
                + (f" — '{got}' is not one of them." if got else " — it is empty.")
            )

        speciality = _text(values.get("speciality"))[:120] or None
        result.data = {"name": name, "email": email, "role": role, "speciality": speciality}
        results.append(result)
    return results


def _dynamic_value(column: Column, raw: Any) -> Any:
    """A spreadsheet cell, turned into what the form would have sent."""
    if _blank(raw):
        return None
    kind = FieldType(column.definition.type) if column.definition else FieldType.TEXT
    if kind == FieldType.BOOLEAN:
        text = _text(raw).lower()
        if isinstance(raw, bool):
            return raw
        if text in {"yes", "y", "true", "1"}:
            return True
        if text in {"no", "n", "false", "0"}:
            return False
        return text  # left for the schema check to name
    if kind == FieldType.MULTISELECT:
        return [part.strip() for part in re.split(r"[;,]", _text(raw)) if part.strip()]
    if kind == FieldType.DATE:
        try:
            parsed = _date(raw)
        except ValueError:
            return _text(raw)
        return parsed.isoformat() if parsed else None
    if kind == FieldType.NUMBER:
        return raw if isinstance(raw, (int, float)) else _text(raw)
    return _text(raw)


async def check_patients(
    session: AsyncSession,
    tenant: Tenant,
    columns: tuple[Column, ...],
    schema: list[FieldDefinition],
    rows: list[tuple[int, dict[str, Any]]],
) -> list[RowResult]:
    """The caller has the clinic bound."""
    existing = {
        (name.lower(), dob)
        for name, dob in (
            await session.execute(
                select(Patient.full_name, Patient.dob).where(Patient.tenant_id == tenant.id)
            )
        ).all()
    }
    extra = [c for c in columns if c.definition is not None]
    seen: dict[tuple[str, date], int] = {}
    results = []
    for number, values in rows:
        name = _text(values.get("full_name"))
        result = RowResult(row=number, label=name or "(no name)")

        try:
            dob = _date(values.get("dob"))
        except ValueError:
            dob = None
            result.errors.append(f"Date of birth '{_text(values.get('dob'))}' is not a date.")
        else:
            if dob is None:
                result.errors.append("Date of birth is required.")

        sex = _text(values.get("sex")).lower() or Sex.OTHER.value
        if sex in {"f"}:
            sex = "female"
        elif sex in {"m"}:
            sex = "male"

        payload = {
            "full_name": name,
            "dob": dob,
            "sex": sex,
            "phone": _text(values.get("phone")),
            "email": _text(values.get("email")) or None,
            "address": _text(values.get("address")) or None,
            "blood_group": _text(values.get("blood_group")) or None,
            "allergies": _text(values.get("allergies")) or None,
            "clinical_ext": {
                c.key: v for c in extra if (v := _dynamic_value(c, values.get(c.key))) is not None
            },
        }

        if dob is not None:
            try:
                PatientIn(**payload)
            except ValidationError as exc:
                for problem in exc.errors():
                    where = ".".join(str(p) for p in problem["loc"])
                    label = next((c.label for c in columns if c.key == where), where)
                    message = problem["msg"].removeprefix("Value error, ")
                    result.errors.append(f"{label}: {message}")
            try:
                payload["clinical_ext"] = tenant_service.validate_dynamic_values(
                    schema, FieldEntity.PATIENT, payload["clinical_ext"]
                )
            except ValidationFailed as exc:
                result.errors += [f["message"] for f in exc.details.get("fields", [])]

            key = (name.lower(), dob)
            if key in existing:
                result.errors.append(f"{name}, born {dob.isoformat()}, is already registered here.")
            elif key in seen:
                result.errors.append(f"The same patient is also on row {seen[key]}.")
            seen.setdefault(key, number)

        result.data = payload
        results.append(result)
    return results
