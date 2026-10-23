"""A clinic's own staff list, in and out of a spreadsheet.

The same sheet the console onboards a clinic with (`app/services/imports.py`),
now in the practice admin's hands: one person per row — name, email, role and
speciality — each opened as an invited account and emailed their own link.
Nobody's password travels in a sheet.

The import only ever adds people. An email the clinic already has is a
duplicate, not an update: changing somebody's role or locking them out is a
decision made on their row of the Team table, one person at a time, not a
side effect of a file. The import runs in the background — see
`app/services/import_jobs.py`.
"""

from __future__ import annotations

import io
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.models.enums import AccountStatus
from app.models.identity import User
from app.models.tenancy import Tenant
from app.services import audit, import_jobs
from app.services.imports import (
    ROLE_LABELS,
    USER_COLUMNS,
    RowResult,
    build_template,
    check_users,
    create_users,
)

KIND = "users"
COLUMNS = USER_COLUMNS


def template(*, clinic: str) -> bytes:
    return build_template(KIND, COLUMNS, clinic=clinic, notes=(
        "An email that already has an account here is a duplicate — the import only adds people.",
        "An export from the Team tab lists everyone, in this same layout.",
    ))


def _row(r: RowResult) -> dict[str, Any]:
    return {
        "row": r.row, "label": r.label, "ok": r.ok, "errors": r.errors,
        "action": "create", "qty": 0, "duplicate_of": r.data.get("duplicate_of"),
    }


async def process(
    session: AsyncSession,
    tenant: Tenant,
    rows: list[tuple[int, dict[str, Any]]],
    *,
    filename: str,
    actor: import_jobs.Actor,
) -> dict[str, Any]:
    """Check every row and open every account, or raise `Rejected` with the
    rows that need fixing. Runs in the worker, in one transaction."""
    results = await check_users(session, tenant, rows)
    invalid = [r for r in results if not r.ok]
    report: dict[str, Any] = {
        "total": len(results), "valid": len(results) - len(invalid), "invalid": len(invalid),
        "duplicates": sum(1 for r in invalid if r.data.get("duplicate_of")),
    }
    if invalid:
        raise import_jobs.Rejected({**report, "rows": [_row(r) for r in invalid]})

    users = await create_users(session, tenant, results)
    created = invited = len(users)
    await audit.record(
        session, action="user.import", tenant_id=tenant.id,
        actor_kind=actor.kind, actor_id=actor.id, actor_label=actor.label,
        entity="import", entity_id=filename,
        meta={"rows": len(results), "created": created, "invited": invited},
    )
    # This runs in the worker, inside the import's transaction, so the accounts
    # are not committed yet: each email job's short countdown, and its retry
    # while the row is not visible, cover the moment until they are.
    from app.worker.tasks import enqueue_link_emails

    enqueue_link_emails(tenant.id, [u.id for u in users], issued_by_label=actor.label)
    return {**report, "created": created, "invited": invited, "rows": []}


def export(users: list[User], *, tz: str) -> bytes:
    """Everyone on the Team tab, in the import's own layout, with the account's
    state after it."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from app.schemas.common import stamp

    book = Workbook()
    sheet = book.active
    sheet.title = "Staff"
    info = ("Account", "Last signed in", "Failed attempts")
    headers = [f"{c.label}{' *' if c.required else ''}" for c in COLUMNS] + list(info)
    sheet.append(headers)
    for i, cell in enumerate(sheet[1]):
        cell.font = Font(bold=True, color="FFFFFF")
        # after the import's own columns, read-only information: an import
        # ignores a header it does not know
        cell.fill = PatternFill("solid", fgColor="4B5563" if i >= len(COLUMNS) else "2F6BF3")

    for user in users:
        sheet.append([
            user.name,
            user.email,
            user.role,
            user.speciality or "",
            AccountStatus(user.status).value,
            stamp(user.last_sign_in, tz) or "",
            user.failed_attempts,
        ])

    for i, header in enumerate(headers, start=1):
        letter = sheet.cell(row=1, column=i).column_letter
        sheet.column_dimensions[letter].width = min(max(len(header) + 4, 16), 40)
    sheet.column_dimensions["B"].width = 34
    sheet.freeze_panes = "B2"

    # the roles, spelt out, for whoever reads the codes in column C
    legend = book.create_sheet("Roles")
    legend.append(["Code", "Role"])
    for code, label in ROLE_LABELS.items():
        legend.append([code, label])
    legend.column_dimensions["A"].width = 16
    legend.column_dimensions["B"].width = 22

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
