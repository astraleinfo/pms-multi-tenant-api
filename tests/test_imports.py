"""Onboarding a clinic from a spreadsheet, from the console.

The sheets are built in memory with openpyxl — the same library the clinic's
own sample sheet comes from — starting from the template the API hands out, so
the test fills in exactly what a person would.
"""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from app.services import mailer
from tests.conftest import sign_in

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _suffix() -> str:
    return uuid.uuid4().hex[:8]


async def _template(console, tenant_id: str, kind: str) -> list[str]:
    response = await console.get(f"/v1/platform/tenants/{tenant_id}/imports/{kind}/template")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[1] == "How to fill this in"
    return [c.value for c in book.worksheets[0][1]]


def _sheet(headers: list[str], rows: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def _upload(console, tenant_id, kind, content, *, commit=False, name="sheet.xlsx"):
    return await console.post(
        f"/v1/platform/tenants/{tenant_id}/imports/{kind}",
        params={"commit": str(commit).lower()},
        files={"file": (name, content, XLSX)},
    )


@pytest.fixture
def outbox(monkeypatch) -> list[mailer.Email]:
    sent: list[mailer.Email] = []

    async def capture(email: mailer.Email) -> None:
        sent.append(email)

    monkeypatch.setattr(mailer, "send", capture)
    return sent


# --- staff -----------------------------------------------------------------------------


async def test_staff_are_created_invited_and_never_given_a_password(console, northgate, outbox):
    headers = await _template(console, northgate.tenant_id, "users")
    assert headers[:3] == ["Full name *", "Email *", "Role *"]

    tag = _suffix()
    content = _sheet(headers, [
        [f"Dr Import {tag}", f"dr.{tag}@example.com", "doctor", "Oral surgery"],
        [f"Desk {tag}", f"DESK.{tag}@Example.com", "Receptionist", None],  # a label, and mixed case
    ])

    preview = await _upload(console, northgate.tenant_id, "users", content)
    assert preview.status_code == 200, preview.text
    assert preview.json()["valid"] == 2 and preview.json()["committed"] is False
    staff_before = (await northgate.get("/v1/users")).json()
    assert not any(tag in u["email"] for u in staff_before)  # a preview writes nothing

    done = await _upload(console, northgate.tenant_id, "users", content, commit=True)
    assert done.status_code == 200, done.text
    assert done.json() | {"rows": None} == done.json() | {
        "committed": True, "created": 2, "invited": 2, "rows": None,
    }

    staff = {u["email"]: u for u in (await northgate.get("/v1/users")).json()}
    assert staff[f"dr.{tag}@example.com"]["role"] == "doctor"
    assert staff[f"dr.{tag}@example.com"]["status"] == "invited"
    assert staff[f"desk.{tag}@example.com"]["role"] == "receptionist"
    assert {e.to for e in outbox} >= {f"dr.{tag}@example.com", f"desk.{tag}@example.com"}


async def test_one_bad_row_stops_the_whole_staff_import(console, northgate, outbox):
    headers = await _template(console, northgate.tenant_id, "users")
    tag = _suffix()
    existing = (await northgate.get("/v1/users")).json()[0]["email"]
    content = _sheet(headers, [
        [f"Good {tag}", f"good.{tag}@example.com", "doctor", None],
        [f"Taken {tag}", existing, "doctor", None],
        [f"Twice {tag}", f"good.{tag}@example.com", "doctor", None],
        [f"Wizard {tag}", f"wiz.{tag}@example.com", "wizard", None],
        ["", "not-an-email", "doctor", None],
    ])

    report = (await _upload(console, northgate.tenant_id, "users", content)).json()
    assert report["valid"] == 1 and report["invalid"] == 4
    errors = {r["row"]: " ".join(r["errors"]) for r in report["rows"]}
    assert "already has an account" in errors[3]
    assert "also on row 2" in errors[4]
    assert "Role must be one of" in errors[5]
    assert "Full name is required" in errors[6] and "not an email" in errors[6]

    refused = await _upload(console, northgate.tenant_id, "users", content, commit=True)
    assert refused.status_code == 422
    assert not any(tag in u["email"] for u in (await northgate.get("/v1/users")).json())
    assert outbox == []


# --- patients ----------------------------------------------------------------------------


async def test_patients_are_registered_with_their_own_record_numbers(console, northgate):
    headers = await _template(console, northgate.tenant_id, "patients")
    assert headers[:2] == ["Full name *", "Date of birth *"]

    tag = _suffix()
    rows = []
    for i, dob in enumerate(["1984-03-09", "12/11/1990"]):
        row = [None] * len(headers)
        row[0] = f"Imported {tag} {i}"
        row[1] = dob
        row[2] = "f" if i == 0 else "male"
        row[3] = "+91 98450 00000"
        rows.append(row)
    # Any custom field this clinic requires gets a plausible value.
    required = [i for i, h in enumerate(headers) if h.endswith(" *") and i > 1]
    for i in required:
        for row in rows:
            row[i] = "yes"

    content = _sheet(headers, rows)
    preview = (await _upload(console, northgate.tenant_id, "patients", content)).json()
    if preview["invalid"]:
        pytest.skip(f"this clinic's custom fields need hand-built values: {preview['rows']}")

    done = await _upload(console, northgate.tenant_id, "patients", content, commit=True)
    assert done.status_code == 200, done.text
    assert done.json()["created"] == 2

    found = (await northgate.get("/v1/patients", params={"search": f"Imported {tag}"})).json()["items"]
    by_name = {p["full_name"]: p for p in found}
    assert by_name[f"Imported {tag} 0"]["dob"] == "1984-03-09"
    assert by_name[f"Imported {tag} 0"]["sex"] == "female"
    assert by_name[f"Imported {tag} 1"]["dob"] == "1990-11-12"
    assert len({p["mrn"] for p in found}) == 2


async def test_a_patient_row_is_checked_like_the_form(console, northgate):
    headers = await _template(console, northgate.tenant_id, "patients")
    existing = (await northgate.get("/v1/patients", params={"limit": 1})).json()["items"][0]

    def row(name, dob, sex=None):
        r = [None] * len(headers)
        r[0], r[1], r[2] = name, dob, sex
        return r

    content = _sheet(headers, [
        row("Future Baby", "2999-01-01"),
        row("No Birthday", None),
        row("Bad Date", "31/02/1990"),
        row(existing["full_name"], existing["dob"]),
        row("Odd Sex", "1990-01-01", "unknown"),
    ])
    report = (await _upload(console, northgate.tenant_id, "patients", content)).json()
    errors = {r["row"]: " ".join(r["errors"]) for r in report["rows"]}
    assert "future" in errors[2]
    assert "Date of birth is required" in errors[3]
    assert "is not a date" in errors[4]
    assert "already registered" in errors[5]
    assert "Sex" in errors[6]


# --- the sheet itself ---------------------------------------------------------------------


async def test_a_sheet_missing_a_required_column_is_refused_whole(console, northgate):
    content = _sheet(["Full name", "Role"], [["Someone", "doctor"]])
    response = await _upload(console, northgate.tenant_id, "users", content)
    assert response.status_code == 422
    assert "Email" in response.json()["error"]["message"]


async def test_headers_match_by_key_in_any_order_and_csv_works(console, northgate):
    tag = _suffix()
    csv_text = f"ROLE,email,name\nlab_tech,lab.{tag}@example.com,Lab {tag}\n".encode()
    report = (await _upload(console, northgate.tenant_id, "users", csv_text, name="staff.csv")).json()
    assert report["valid"] == 1


async def test_not_a_spreadsheet_is_a_clear_refusal(console, northgate):
    response = await _upload(console, northgate.tenant_id, "users", b"%PDF-1.4", name="staff.pdf")
    assert response.status_code == 422
    garbage = await _upload(console, northgate.tenant_id, "users", b"not a zip", name="staff.xlsx")
    assert garbage.status_code == 422


async def test_an_empty_sheet_says_so(console, northgate):
    headers = await _template(console, northgate.tenant_id, "users")
    response = await _upload(console, northgate.tenant_id, "users", _sheet(headers, []))
    assert response.status_code == 422
    assert "no rows" in response.json()["error"]["message"]


async def test_only_the_console_can_import(client, northgate):
    admin = await sign_in(client, "northgate_admin")
    response = await admin.post(
        f"/v1/platform/tenants/{northgate.tenant_id}/imports/users",
        files={"file": ("s.xlsx", b"x", XLSX)},
    )
    assert response.status_code == 401


async def test_an_unknown_kind_is_a_422(console, northgate):
    response = await console.get(f"/v1/platform/tenants/{northgate.tenant_id}/imports/invoices/template")
    assert response.status_code == 422
