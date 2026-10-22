"""A clinic's staff list in and out of a spreadsheet, from the Team tab.

As with stock, the Celery hop is replaced by a recorder and the worker half
(`import_jobs.run`) is called directly. Invitation emails are captured, never
sent.
"""

from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import Workbook, load_workbook

from app.db.session import session_factory, tenant_session
from app.models.imports import ImportJob
from app.services import import_jobs, mailer
from app.worker import tasks

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _tag() -> str:
    return uuid.uuid4().hex[:8]


@pytest.fixture(autouse=True)
def queued(monkeypatch) -> list[tuple[uuid.UUID, uuid.UUID]]:
    calls: list[tuple[uuid.UUID, uuid.UUID]] = []
    monkeypatch.setattr(tasks, "enqueue_import", lambda t, j: calls.append((t, j)) or True)
    return calls


@pytest.fixture
def outbox(monkeypatch) -> list[mailer.Email]:
    sent: list[mailer.Email] = []

    async def capture(email: mailer.Email) -> None:
        sent.append(email)

    monkeypatch.setattr(mailer, "send", capture)
    return sent


async def _headers(clinic) -> list[str]:
    response = await clinic.get("/v1/users/import/template")
    assert response.status_code == 200, response.text
    return [c.value for c in load_workbook(io.BytesIO(response.content)).worksheets[0][1]]


def _sheet(headers: list[str], rows: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def _import(clinic, content: bytes) -> dict:
    queued = await clinic.post("/v1/users/import", files={"file": ("staff.xlsx", content, XLSX)})
    assert queued.status_code == 202, queued.text
    job_id = queued.json()["id"]
    await import_jobs.run(session_factory, uuid.UUID(clinic.tenant_id), uuid.UUID(job_id))
    job = await clinic.get(f"/v1/users/import/jobs/{job_id}")
    assert job.status_code == 200, job.text
    return job.json()


async def _staff(clinic) -> dict[str, dict]:
    return {u["email"]: u for u in (await clinic.get("/v1/users")).json()}


async def test_the_sample_sheet_is_the_staff_sheet(northgate):
    headers = await _headers(northgate)
    assert headers == ["Full name *", "Email *", "Role *", "Speciality"]


async def test_people_are_opened_invited_and_emailed_in_the_background(northgate, outbox, link_emails):
    tag = _tag()
    content = _sheet(await _headers(northgate), [
        [f"Dr Team {tag}", f"dr.{tag}@example.com", "doctor", "Periodontics"],
        [f"Desk {tag}", f"desk.{tag}@example.com", "Receptionist", None],
    ])
    job = await _import(northgate, content)
    assert job["status"] == "done", job
    assert (job["report"]["created"], job["report"]["invited"]) == (2, 2)

    # The invitations were queued from inside the import's transaction; each
    # is its own email job, run once the accounts exist.
    assert not outbox
    assert await link_emails.deliver() == ["sent", "sent"]

    staff = await _staff(northgate)
    assert staff[f"dr.{tag}@example.com"]["status"] == "invited"
    assert staff[f"dr.{tag}@example.com"]["role"] == "doctor"
    assert staff[f"desk.{tag}@example.com"]["role"] == "receptionist"
    assert {e.to for e in outbox} >= {f"dr.{tag}@example.com", f"desk.{tag}@example.com"}


async def test_an_email_already_here_is_a_duplicate_and_nobody_is_added(northgate, outbox):
    tag = _tag()
    existing = next(iter(await _staff(northgate)))
    content = _sheet(await _headers(northgate), [
        [f"New {tag}", f"new.{tag}@example.com", "pharmacist", None],
        ["Someone", existing, "doctor", None],
        [f"Twice {tag}", f"new.{tag}@example.com", "doctor", None],
        [f"No role {tag}", f"norole.{tag}@example.com", "wizard", None],
    ])
    job = await _import(northgate, content)
    assert job["status"] == "failed"
    assert job["error"] == "2 duplicate(s) and 1 other row(s) to fix — nothing was imported."
    dupes = {r["row"]: r["duplicate_of"] for r in job["report"]["rows"] if r["duplicate_of"]}
    assert dupes == {3: existing, 4: "row 2"}
    assert f"new.{tag}@example.com" not in await _staff(northgate), "all or nothing"
    assert outbox == []


async def test_the_export_lists_everyone_in_the_import_layout(northgate):
    response = await northgate.get("/v1/users/export")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Staff", "Roles"]
    head = [c.value for c in book["Staff"][1]]
    assert head[:4] == await _headers(northgate)
    emails = {row[1] for row in book["Staff"].iter_rows(min_row=2, values_only=True)}
    assert emails == set(await _staff(northgate))


async def test_a_file_that_is_not_the_staff_sheet_is_refused_at_the_upload(northgate, queued):
    before = len(queued)
    stock_sheet = _sheet(["SKU *", "Item name *", "Category *"], [["X", "X", "lens"]])
    response = await northgate.post("/v1/users/import", files={"file": ("staff.xlsx", stock_sheet, XLSX)})
    assert response.status_code == 422
    assert "missing required columns" in response.json()["error"]["message"]
    assert len(queued) == before


async def test_a_staff_import_job_is_not_reachable_as_a_stock_one(northgate):
    content = _sheet(await _headers(northgate), [[f"Kind {_tag()}", f"kind.{_tag()}@example.com", "doctor", None]])
    job_id = (await northgate.post("/v1/users/import", files={"file": ("staff.xlsx", content, XLSX)})).json()["id"]
    assert (await northgate.get(f"/v1/inventory/import/jobs/{job_id}")).status_code == 404
    assert (await northgate.get("/v1/users/import/jobs/latest")).json()["id"] == job_id
    await import_jobs.run(session_factory, uuid.UUID(northgate.tenant_id), uuid.UUID(job_id))


async def test_a_lost_hand_off_is_retried_by_hand(northgate, queued, outbox):
    """No sweep: a job the broker never got stays queued, turns retryable once it
    has sat a while, and the person hands it over again."""
    tag = _tag()
    content = _sheet(await _headers(northgate), [[f"Lost {tag}", f"lost.{tag}@example.com", "doctor", None]])
    job = (await northgate.post("/v1/users/import", files={"file": ("staff.xlsx", content, XLSX)})).json()
    assert job["status"] == "queued" and job["retryable"] is False
    retry_url = f"/v1/users/import/jobs/{job['id']}/retry"
    # Too soon: its hand-off may still be on the way.
    assert (await northgate.post(retry_url)).status_code == 409

    tenant_id, job_id = uuid.UUID(northgate.tenant_id), uuid.UUID(job["id"])
    async with tenant_session(tenant_id) as session:
        row = await session.get(ImportJob, job_id)
        row.created_at = row.updated_at = row.created_at - import_jobs.STUCK_AFTER * 2
    assert (await northgate.get(f"/v1/users/import/jobs/{job['id']}")).json()["retryable"] is True
    # A staff job is not reachable through the stock route, retry included.
    assert (await northgate.post(f"/v1/inventory/import/jobs/{job['id']}/retry")).status_code == 404

    queued.clear()
    retried = await northgate.post(retry_url)
    assert retried.status_code == 200, retried.text
    assert retried.json()["retryable"] is False
    assert queued == [(tenant_id, job_id)]

    await import_jobs.run(session_factory, tenant_id, job_id)
    assert (await northgate.post(retry_url)).status_code == 409


async def test_only_an_admin_can_import_or_export_staff(northgate_doctor):
    assert (await northgate_doctor.get("/v1/users/export")).status_code == 403
    assert (await northgate_doctor.get("/v1/users/import/template")).status_code == 403
    content = _sheet(["Full name *", "Email *", "Role *"], [["X", "x@example.com", "doctor"]])
    response = await northgate_doctor.post("/v1/users/import", files={"file": ("s.xlsx", content, XLSX)})
    assert response.status_code == 403
