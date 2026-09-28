"""Stock lines in and out of a spreadsheet, from the Inventory page.

Sheets are built with openpyxl from the headers of the template the API hands
out, so each test fills in what a person would.

The Celery hop is not exercised: `enqueue_import` is replaced with a
recorder, and the worker half (`import_jobs.run`) is called directly against
the test database — which is everything the task does.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.db.session import session_factory, tenant_session
from app.models.imports import ImportJob
from app.services import import_jobs
from app.worker import tasks

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _tag() -> str:
    return uuid.uuid4().hex[:6].upper()


async def _headers(clinic) -> list[str]:
    response = await clinic.get("/v1/inventory/import/template")
    assert response.status_code == 200, response.text
    return [c.value for c in load_workbook(io.BytesIO(response.content)).worksheets[0][1]]


def _sheet(headers: list[str], rows: list[dict], *, unique_names: bool = True) -> bytes:
    """Rows keyed by header label without the required-marker.

    The test database outlives a run, and a name already in stock under
    another SKU is a duplicate — so by default each item's name carries its
    own SKU, and a test about duplicates asks for the names as written."""
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    plain = [h.removesuffix(" *") for h in headers]
    for row in rows:
        if unique_names and row.get("Item name") and row.get("SKU"):
            row = {**row, "Item name": f"{row['Item name']} {row['SKU']}"}
        sheet.append([row.get(h) for h in plain])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture(autouse=True)
def queued(monkeypatch) -> list[tuple[uuid.UUID, uuid.UUID]]:
    calls: list[tuple[uuid.UUID, uuid.UUID]] = []
    monkeypatch.setattr(tasks, "enqueue_import", lambda t, j: calls.append((t, j)) or True)
    return calls


async def _upload(clinic, content: bytes, *, name="stock.xlsx"):
    return await clinic.post("/v1/inventory/import", files={"file": (name, content, XLSX)})


async def _import(clinic, content: bytes, *, name="stock.xlsx") -> dict:
    """Upload, run the job as the worker would, and read back how it went."""
    queued = await _upload(clinic, content, name=name)
    assert queued.status_code == 202, queued.text
    assert queued.json()["status"] == "queued"
    job_id = queued.json()["id"]
    await import_jobs.run(session_factory, uuid.UUID(clinic.tenant_id), uuid.UUID(job_id))
    job = await clinic.get(f"/v1/inventory/import/jobs/{job_id}")
    assert job.status_code == 200, job.text
    return job.json()


async def _stored_file(tenant_id: str, job_id: str):
    async with tenant_session(uuid.UUID(tenant_id)) as session:
        return (
            await session.execute(
                select(ImportJob.data, ImportJob.storage_key).where(ImportJob.id == uuid.UUID(job_id))
            )
        ).one()


async def _item(clinic, sku: str) -> dict | None:
    return next((i for i in (await clinic.get("/v1/inventory/items")).json() if i["sku"] == sku), None)


async def _on_hand(clinic, item_id: str) -> int:
    levels = (await clinic.get("/v1/inventory/levels")).json()
    return next(level["on_hand"] for level in levels if level["item_id"] == item_id)


def _later(days: int = 400) -> str:
    return (date.today() + timedelta(days=days)).isoformat()


async def test_the_sample_sheet_carries_every_field_and_an_example_per_pricing_mode(northgate):
    response = await northgate.get("/v1/inventory/import/template")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Items", "How to fill this in"]

    sheet = book.worksheets[0]
    headers = [c.value for c in sheet[1]]
    assert headers[:3] == ["SKU *", "Item name *", "Category *"]
    for label in ("Pricing mode", "Selling price", "Cost price", "MRP", "Margin %", "HSN code",
                  "Quantity received", "Batch number", "Expiry date"):
        assert label in headers
    modes = {sheet.cell(row=r, column=headers.index("Pricing mode") + 1).value for r in (2, 3, 4)}
    assert modes == {"price", "percent", "mrp"}
    listed = {v.formula1 for v in sheet.data_validations.dataValidation}
    assert '"medicine,consumable,implant,lens,reagent"' in listed
    assert '"price,percent,mrp"' in listed


async def test_the_sample_sheet_imports_as_it_comes(northgate):
    """The three example rows are valid rows — a person who changes only the
    codes has a working sheet."""
    response = await northgate.get("/v1/inventory/import/template")
    book = load_workbook(io.BytesIO(response.content))
    sheet = book.worksheets[0]
    tag = _tag()
    for r in (2, 3, 4):
        sheet.cell(row=r, column=1).value = f"{sheet.cell(row=r, column=1).value}-{tag}"
        sheet.cell(row=r, column=2).value = f"{sheet.cell(row=r, column=2).value} {tag}"
    out = io.BytesIO()
    book.save(out)

    job = await _import(northgate, out.getvalue())
    assert job["status"] == "done", job
    assert job["report"]["created"] == 3 and job["report"]["stocked"] == 3


async def test_new_lines_are_opened_priced_and_stocked_as_a_purchase(northgate):
    headers = await _headers(northgate)
    sku = f"IMP-{_tag()}"
    content = _sheet(headers, [{
        "SKU": sku.lower(), "Item name": "Imported amoxicillin", "Category": "Medicine",
        "Form": "Capsule", "Strength": "500 mg", "Unit": "capsule", "Reorder level": 30,
        "Pricing mode": "percent", "Cost price": 12, "Margin %": 25, "HSN code": "3004",
        "Quantity received": 120, "Batch number": "B-1", "Expiry date": _later(),
    }])

    queued = await _upload(northgate, content)
    assert queued.status_code == 202, queued.text
    assert await _item(northgate, sku) is None, "nothing is written until the worker runs"

    job_id = queued.json()["id"]
    assert await import_jobs.run(session_factory, uuid.UUID(northgate.tenant_id), uuid.UUID(job_id)) == "done"
    job = (await northgate.get(f"/v1/inventory/import/jobs/{job_id}")).json()
    assert job["status"] == "done" and job["error"] == "" and job["finished_at"]
    report = job["report"]
    assert report["committed"] is True
    assert (report["created"], report["updated"], report["stocked"]) == (1, 0, 1)

    item = await _item(northgate, sku)
    assert item is not None, "the SKU is stored in capitals"
    assert item["category"] == "medicine" and item["pricing_mode"] == "percent"
    assert item["unit_price"] == 15.0, "12.00 at 25% is 15.00 — the server did the arithmetic"
    assert await _on_hand(northgate, item["id"]) == 120

    ledger = (await northgate.get("/v1/inventory/transactions", params={"item_id": item["id"]})).json()
    assert [(t["type"], t["qty"], t["batch_no"]) for t in ledger] == [("purchase", 120, "B-1")]
    assert "stock.xlsx" in ledger[0]["note"]


async def test_a_known_sku_updates_and_blank_cells_leave_fields_alone(northgate):
    sku = f"UPD-{_tag()}"
    created = await northgate.post("/v1/inventory/items", json={
        "sku": sku, "name": "Composite resin", "category": "consumable", "unit_price": 900,
        "manufacturer": "3M", "reorder_level": 4,
    })
    assert created.status_code == 201, created.text

    headers = await _headers(northgate)
    content = _sheet(headers, [{"SKU": sku, "Selling price": 950, "Reorder level": 6}])
    job = await _import(northgate, content)
    assert job["status"] == "done", job
    assert (job["report"]["created"], job["report"]["updated"], job["report"]["stocked"]) == (0, 1, 0)
    item = await _item(northgate, sku)
    assert item["unit_price"] == 950.0 and item["reorder_level"] == 6
    assert item["name"] == "Composite resin" and item["manufacturer"] == "3M"
    assert item["category"] == "consumable"


async def test_one_bad_row_holds_back_the_whole_sheet_and_each_says_why(northgate):
    headers = await _headers(northgate)
    tag = _tag()
    good = {"SKU": f"OK-{tag}", "Item name": "Fine", "Category": "consumable"}
    content = _sheet(headers, [
        good,
        {"SKU": f"NOCAT-{tag}", "Item name": "No category"},
        {"SKU": f"MODE-{tag}", "Item name": "Bad mode", "Category": "medicine", "Pricing mode": "cheap"},
        {"SKU": f"PCT-{tag}", "Item name": "No cost", "Category": "medicine", "Pricing mode": "percent",
         "Margin %": 20},
        {"SKU": f"QTY-{tag}", "Item name": "No batch", "Category": "medicine", "Quantity received": 10},
        {"SKU": f"OLD-{tag}", "Item name": "Expired", "Category": "medicine", "Quantity received": 5,
         "Batch number": "X", "Expiry date": "2020-01-01"},
        {"SKU": f"NEG-{tag}", "Item name": "Negative", "Category": "medicine", "Quantity received": -3,
         "Batch number": "X", "Expiry date": _later()},
        {"SKU": f"PRICE-{tag}", "Item name": "Word price", "Category": "medicine", "Selling price": "lots"},
        {**good, "Item name": "Twice"},
    ])

    job = await _import(northgate, content)
    assert job["status"] == "failed", job
    assert job["error"] == "1 duplicate(s) and 7 other row(s) to fix — nothing was imported."
    body = job["report"]
    assert (body["valid"], body["invalid"], body["duplicates"]) == (1, 8, 1), body
    assert body["committed"] is False
    assert all(not r["ok"] for r in body["rows"]), "only the rows to fix come back"
    errors = {r["row"]: " ".join(r["errors"]) for r in body["rows"]}
    assert "Category is required" in errors[3]
    assert "Pricing mode must be one of" in errors[4]
    assert "needs a cost price" in errors[5]
    assert "Batch number is needed" in errors[6] and "Expiry date is needed" in errors[6]
    assert "has already passed" in errors[7]
    assert "cannot be negative" in errors[8]
    assert "is not a number" in errors[9]
    assert "also on row 2" in errors[10]
    assert next(r for r in body["rows"] if r["row"] == 10)["duplicate_of"] == "row 2"

    assert await _item(northgate, f"OK-{tag}") is None, "all or nothing"


async def test_a_batch_already_on_the_shelf_must_keep_its_expiry(northgate):
    headers = await _headers(northgate)
    sku = f"BAT-{_tag()}"
    first = _sheet(headers, [{"SKU": sku, "Item name": "Lidocaine", "Category": "medicine",
                              "Quantity received": 10, "Batch number": "L1", "Expiry date": _later(300)}])
    assert (await _import(northgate, first))["status"] == "done"

    again = _sheet(headers, [{"SKU": sku, "Quantity received": 5, "Batch number": "L1",
                              "Expiry date": _later(300)}])
    clash = _sheet(headers, [{"SKU": sku, "Quantity received": 5, "Batch number": "L1",
                              "Expiry date": _later(500)}])
    failed = await _import(northgate, clash)
    assert failed["status"] == "failed"
    assert "already on the shelf" in " ".join(failed["report"]["rows"][0]["errors"])

    assert (await _import(northgate, again))["status"] == "done"
    assert await _on_hand(northgate, (await _item(northgate, sku))["id"]) == 15


async def test_the_export_goes_back_in_as_an_update_that_adds_no_stock(northgate):
    headers = await _headers(northgate)
    sku = f"EXP-{_tag()}"
    seed = _sheet(headers, [{"SKU": sku, "Item name": "Export check", "Category": "medicine",
                             "Pricing mode": "mrp", "MRP": 50, "Margin %": 10,
                             "Quantity received": 40, "Batch number": "E1", "Expiry date": _later()}])
    assert (await _import(northgate, seed))["status"] == "done"

    response = await northgate.get("/v1/inventory/export")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX)
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Items", "Batches"]

    items = book["Items"]
    head = [c.value for c in items[1]]
    assert head[: len(headers)] == headers, "the import's own layout, first"
    row = next(r for r in items.iter_rows(min_row=2, values_only=True) if r[0] == sku)
    assert row[head.index("On hand")] == 40
    assert row[head.index("Selling price")] == 45.0
    assert row[head.index("Quantity received")] is None

    batches = [r for r in book["Batches"].iter_rows(min_row=2, values_only=True) if r[0] == sku]
    assert [(r[2], r[4]) for r in batches] == [("E1", 40)]

    back = await _import(northgate, response.content, name="export.xlsx")
    assert back["status"] == "done", back
    assert back["report"]["created"] == 0 and back["report"]["stocked"] == 0
    assert await _on_hand(northgate, (await _item(northgate, sku))["id"]) == 40, "no stock added"


async def test_only_people_who_can_change_stock_can_import(northgate_doctor):
    content = _sheet(["SKU *", "Item name *", "Category *"], [{"SKU": "X", "Item name": "X", "Category": "lens"}])
    assert (await _upload(northgate_doctor, content)).status_code == 403
    assert (await northgate_doctor.get("/v1/inventory/import/template")).status_code == 403


async def test_the_upload_is_deleted_when_the_job_ends_either_way(northgate, queued):
    headers = await _headers(northgate)
    good = _sheet(headers, [{"SKU": f"DEL-{_tag()}", "Item name": "Kept briefly", "Category": "lens"}])
    bad = _sheet(headers, [{"SKU": f"DEL-{_tag()}", "Item name": "No category"}])

    for content, outcome in ((good, "done"), (bad, "failed")):
        job_id = (await _upload(northgate, content)).json()["id"]
        assert queued[-1] == (uuid.UUID(northgate.tenant_id), uuid.UUID(job_id)), "handed to the worker"
        data, key = await _stored_file(northgate.tenant_id, job_id)
        assert data == content and key is None, "held on the row until the worker has it"

        tenant_id, job = uuid.UUID(northgate.tenant_id), uuid.UUID(job_id)
        assert await import_jobs.run(session_factory, tenant_id, job) == outcome
        assert await _stored_file(northgate.tenant_id, job_id) == (None, None)
        # A redelivered job — the worker restarted after finishing — does nothing.
        assert await import_jobs.run(session_factory, tenant_id, job) == "skipped"


async def test_a_file_that_is_not_the_sheet_is_refused_at_the_upload(northgate, queued):
    """Checked before a job exists: the person is answered at once, and
    nothing is stored or handed to the worker."""
    headers = await _headers(northgate)
    cases = [
        ("notes.pdf", b"%PDF-1.7", "Upload the .xlsx sample sheet"),
        ("stock.xlsx", b"", "empty"),
        ("stock.xlsx", b"not a workbook at all", "could not be read as an .xlsx"),
        ("stock.xlsx", _sheet(["Patient", "Date of birth"], [{"Patient": "x"}]), "missing required columns"),
        ("stock.xlsx", _sheet(headers, []), "no rows"),
        ("stock.csv", "SKU,Item name,Category\n".encode("utf-16"), "not UTF-8"),
    ]
    before = len(queued)
    for name, content, says in cases:
        response = await _upload(northgate, content, name=name)
        assert response.status_code == 422, (name, response.text)
        assert says in response.json()["error"]["message"], (name, response.json())
    assert len(queued) == before


async def test_a_new_code_for_an_item_already_in_stock_is_a_duplicate(northgate):
    headers = await _headers(northgate)
    tag = _tag()
    name = f"Chlorhexidine {tag}"
    kept = _sheet(headers, [{"SKU": f"CHX-{tag}", "Item name": name, "Category": "medicine",
                             "Form": "Mouthwash", "Strength": "0.2%"}], unique_names=False)
    assert (await _import(northgate, kept))["status"] == "done"

    again = _sheet(headers, [
        # the same product, spelt with different case and spacing, under a new code
        {"SKU": f"CHX2-{tag}", "Item name": f"  chlorhexidine   {tag.lower()}", "Category": "medicine",
         "Form": "mouthwash", "Strength": "0.2%"},
        # the same product twice in the sheet, under two new codes
        {"SKU": f"GEL-{tag}", "Item name": f"Gel {tag}", "Category": "medicine", "Strength": "1%"},
        {"SKU": f"GEL2-{tag}", "Item name": f"Gel {tag}", "Category": "medicine", "Strength": "1%"},
        # a different strength is a different product
        {"SKU": f"GEL5-{tag}", "Item name": f"Gel {tag}", "Category": "medicine", "Strength": "5%"},
        # its own SKU is an update, not a duplicate
        {"SKU": f"CHX-{tag}", "Reorder level": 3},
    ], unique_names=False)
    job = await _import(northgate, again)
    assert job["status"] == "failed"
    assert job["error"] == "2 duplicate(s) — nothing was imported."
    report = job["report"]
    assert (report["invalid"], report["duplicates"]) == (2, 2)
    dupes = {r["row"]: r["duplicate_of"] for r in report["rows"]}
    assert dupes == {2: f"CHX-{tag}", 4: "row 3"}
    assert await _item(northgate, f"GEL-{tag}") is None, "all or nothing"


async def test_the_latest_job_is_there_for_a_page_opened_mid_import(northgate):
    job_id = (await _upload(northgate, _sheet(await _headers(northgate), [
        {"SKU": f"LAT-{_tag()}", "Item name": "Latest", "Category": "reagent"},
    ]))).json()["id"]
    latest = (await northgate.get("/v1/inventory/import/jobs/latest")).json()
    assert latest["id"] == job_id and latest["status"] == "queued"
    await import_jobs.run(session_factory, uuid.UUID(northgate.tenant_id), uuid.UUID(job_id))
